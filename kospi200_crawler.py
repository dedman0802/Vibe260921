"""네이버 증권(https://stock.naver.com/market/stock/kr)의 코스피200 데이터를 수집한다.

이 페이지는 Next.js로 만든 화면이라 HTML을 내려받으면 코스피200 수치가 없고(BeautifulSoup으로
파싱할 내용이 없음), 브라우저에서 JavaScript가 JSON API를 호출해 화면을 채운다.
그래서 HTML을 파싱하는 대신 페이지가 실제로 쓰는 JSON API를 직접 호출한다.

    지수 현재값   https://polling.finance.naver.com/api/realtime/domestic/index/KPI200
    일별 시세     https://m.stock.naver.com/api/index/KPI200/price
    구성 종목     https://m.stock.naver.com/api/index/KPI200/enrollStocks   (한 번에 최대 50건)

사용 예:
    python kospi200_crawler.py
    python kospi200_crawler.py --days 60 --out kospi200.xlsx
    python kospi200_crawler.py --top 20        (화면에는 시가총액 상위 20개만 출력)
    python kospi200_crawler.py --out kospi200.json

필요 패키지: requests, openpyxl(엑셀 저장 시)
"""
import argparse
import json
import sys
import time

import requests

INDEX_CODE = "KPI200"
POLLING_URL = f"https://polling.finance.naver.com/api/realtime/domestic/index/{INDEX_CODE}"
PRICE_URL = f"https://m.stock.naver.com/api/index/{INDEX_CODE}/price"
STOCKS_URL = f"https://m.stock.naver.com/api/index/{INDEX_CODE}/enrollStocks"
PAGE_SIZE = 50  # 서버가 허용하는 최대값 (100 이상이면 400 오류)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Referer": "https://stock.naver.com/",
}


def get_json(url, **params):
    res = requests.get(url, params=params, headers=HEADERS, timeout=10)
    res.raise_for_status()
    return res.json()


def num(text):
    """'1,112.16', '92,658천주' 같은 문자열에서 숫자만 뽑아 int/float로 바꾼다. 값이 없으면 None."""
    if text is None:
        return None
    cleaned = "".join(ch for ch in str(text) if ch.isdigit() or ch in ".-")
    if cleaned in ("", "-", "."):
        return None
    value = float(cleaned)
    return int(value) if value.is_integer() and "." not in cleaned else value


def signed(item, key):
    """전일 대비 값에 등락 방향을 반영한다. (API는 하락일 때도 절댓값을 준다)"""
    value = num(item.get(key))
    if value is None:
        return None
    falling = (item.get("compareToPreviousPrice") or {}).get("name") == "FALLING"
    return -abs(value) if falling else value


def fetch_summary():
    """코스피200 지수의 현재 시세 요약."""
    data = get_json(POLLING_URL)["datas"][0]
    return {
        "name": data["stockName"],
        "code": data["itemCode"],
        "close": num(data["closePrice"]),
        "change": signed(data, "compareToPreviousClosePrice"),
        "change_rate": signed(data, "fluctuationsRatio"),
        "open": num(data["openPrice"]),
        "high": num(data["highPrice"]),
        "low": num(data["lowPrice"]),
        "volume": num(data["accumulatedTradingVolumeRaw"]),  # 주
        "trading_value": num(data["accumulatedTradingValueRaw"]),  # 원
        "market_status": data["marketStatus"],  # OPEN / CLOSE
        "traded_at": data["localTradedAt"],
    }


def fetch_history(days):
    """최근 일별 시세 (최신 날짜가 먼저)."""
    rows, page = [], 1
    while len(rows) < days:
        chunk = get_json(PRICE_URL, page=page, pageSize=min(PAGE_SIZE, days - len(rows)))
        if not chunk:
            break
        rows += chunk
        page += 1
    return [
        {
            "date": r["localTradedAt"],
            "close": num(r["closePrice"]),
            "change": signed(r, "compareToPreviousClosePrice"),
            "change_rate": signed(r, "fluctuationsRatio"),
            "open": num(r["openPrice"]),
            "high": num(r["highPrice"]),
            "low": num(r["lowPrice"]),
        }
        for r in rows[:days]
    ]


def fetch_constituents(delay=0.2):
    """코스피200 구성 종목 전체 (약 200개). 페이지를 넘기며 빈 응답이 나올 때까지 요청한다."""
    stocks, seen, page = [], set(), 1
    while True:
        chunk = get_json(STOCKS_URL, page=page, pageSize=PAGE_SIZE)
        new = [s for s in chunk if s["itemCode"] not in seen]
        if not new:  # 빈 페이지(또는 이미 받은 종목만 있는 페이지)면 마지막 페이지까지 다 받은 것
            break
        print(f"  {page}페이지: {len(new)}건 (누적 {len(stocks) + len(new)}건)")
        for s in new:
            seen.add(s["itemCode"])
            stocks.append({
                "code": s["itemCode"],
                "name": s["stockName"],
                "price": num(s["closePrice"]),
                "change": signed(s, "compareToPreviousClosePrice"),
                "change_rate": signed(s, "fluctuationsRatio"),
                "volume": num(s.get("accumulatedTradingVolume")),  # 주
                "trading_value": num(s.get("accumulatedTradingValue")),  # 백만원
                "market_cap": num(s.get("marketValue")),  # 억원
            })
        page += 1
        time.sleep(delay)  # 서버에 부담을 주지 않도록 요청 사이에 간격을 둔다
    return stocks


SHEETS = [
    # (시트 이름, 결과 키, [(머리글, 필드), ...])
    ("지수 요약", "summary", [
        ("지수명", "name"), ("코드", "code"), ("현재값", "close"), ("전일대비", "change"),
        ("등락률(%)", "change_rate"), ("시가", "open"), ("고가", "high"), ("저가", "low"),
        ("거래량(주)", "volume"), ("거래대금(원)", "trading_value"), ("장 상태", "market_status"),
        ("기준 시각", "traded_at"),
    ]),
    ("일별 시세", "history", [
        ("날짜", "date"), ("종가", "close"), ("전일대비", "change"), ("등락률(%)", "change_rate"),
        ("시가", "open"), ("고가", "high"), ("저가", "low"),
    ]),
    ("구성 종목", "constituents", [
        ("종목코드", "code"), ("종목명", "name"), ("현재가", "price"), ("전일대비", "change"),
        ("등락률(%)", "change_rate"), ("거래량(주)", "volume"), ("거래대금(백만원)", "trading_value"),
        ("시가총액(억원)", "market_cap"),
    ]),
]


def save_excel(result, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for title, key, columns in SHEETS:
        ws = wb.create_sheet(title)
        rows = result[key] if isinstance(result[key], list) else [result[key]]
        ws.append([header for header, _ in columns])
        for row in rows:
            ws.append([row.get(field) for _, field in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.alignment = Alignment(horizontal="center")
        for i, (header, field) in enumerate(columns, 1):
            width = max(len(header) * 2, *(len(str(r.get(field, ""))) for r in rows)) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cell:
                    if isinstance(c.value, (int, float)):
                        c.number_format = "#,##0.00" if isinstance(c.value, float) else "#,##0"
        ws.freeze_panes = "A2"
    wb.save(path)


def save(result, path):
    if path.lower().endswith(".xlsx"):
        save_excel(result, path)
    else:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)


def fmt(value, digits=2):
    return "-" if value is None else f"{value:,.{digits}f}" if isinstance(value, float) else f"{value:,}"


def print_report(result, top_n=0):
    s = result["summary"]
    print(f"\n■ {s['name']} ({s['code']})  [{s['market_status']}, {s['traded_at']}]")
    print(f"  현재값 {fmt(s['close'])}  전일대비 {fmt(s['change'])} ({fmt(s['change_rate'])}%)")
    print(f"  시가 {fmt(s['open'])}  고가 {fmt(s['high'])}  저가 {fmt(s['low'])}")
    print(f"  거래량 {fmt(s['volume'])}주  거래대금 {fmt(s['trading_value'])}원")

    print(f"\n■ 최근 일별 시세 ({len(result['history'])}일 중 5일)")
    for r in result["history"][:5]:
        print(f"  {r['date']}  종가 {fmt(r['close']):>10}  {fmt(r['change']):>8} ({fmt(r['change_rate'])}%)")

    ranked = sorted(result["constituents"], key=lambda x: x["market_cap"] or 0, reverse=True)
    top = ranked[:top_n] if top_n > 0 else ranked  # 0이면 전체 출력
    scope = f"시가총액 상위 {top_n}개" if top_n > 0 else "시가총액 순 전체"
    print(f"\n■ 구성 종목 {len(ranked)}개 ({scope})")
    for rank, r in enumerate(top, 1):
        print(f"  {rank:>2}. {r['name']:<12} {fmt(r['price']):>10}원  {fmt(r['change_rate']):>6}%  시총 {fmt(r['market_cap']):>12}억")


def main():
    parser = argparse.ArgumentParser(description="네이버 증권 코스피200 크롤러")
    parser.add_argument("--days", type=int, default=30, help="가져올 일별 시세 일수 (기본 30)")
    parser.add_argument("--top", type=int, default=0, help="화면에 출력할 시가총액 상위 종목 수 (기본 0 = 전체)")
    parser.add_argument("--out", default="kospi200.xlsx", help="저장할 파일 경로 (.xlsx 또는 .json)")
    args = parser.parse_args()

    try:
        print("지수 요약 수집 중...")
        summary = fetch_summary()
        print("일별 시세 수집 중...")
        history = fetch_history(args.days)
        print("구성 종목 수집 중...")
        constituents = fetch_constituents()
    except (requests.RequestException, KeyError, IndexError, ValueError) as e:
        sys.exit(f"수집 실패: {e}")

    result = {"summary": summary, "history": history, "constituents": constituents}
    print_report(result, args.top)

    try:
        save(result, args.out)
    except OSError as e:  # 같은 이름의 파일이 엑셀에서 열려 있는 경우 등
        sys.exit(f"저장 실패: {e}")
    print(f"\n{args.out}에 저장했습니다.")


if __name__ == "__main__":
    main()

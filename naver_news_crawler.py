"""네이버 검색 결과(뉴스)에서 기사 목록을 수집하고 기사 본문을 크롤링한다.

검색 결과 HTML 구조(2026-09 기준):
    기사 하나는 같은 gdid(data-nlog-params)를 가진 여러 <a> 태그로 이루어진다.
        a[data-heatmap-target=".tit"]   제목 (제목 안의 <mark>는 검색어 강조)
        a[data-heatmap-target=".body"]  요약 (대표 기사에만 있음)
        a[data-heatmap-target=".img"]   썸네일 (대표 기사에만 있음)
        a[data-heatmap-target=".nav"]   '네이버뉴스' 링크 -> n.news.naver.com 기사 주소
        [data-sds-comp="Profile"]       언론사명(.sds-comps-profile-info-title-text)과
                                        게시 시각(.sds-comps-profile-info-subtext, 예: "7시간 전")

사용 예:
    python naver_news_crawler.py
    python naver_news_crawler.py --limit 5 --out news.json
    python naver_news_crawler.py --limit 5 --out news.xlsx   (확장자가 .xlsx면 엑셀로 저장)
    python naver_news_crawler.py --url "<네이버 검색 URL>"

필요 패키지: requests, beautifulsoup4, openpyxl(엑셀 저장 시)
"""
import argparse
import json
import re
import sys
import time

import requests
from bs4 import BeautifulSoup

DEFAULT_URL = (
    "https://search.naver.com/search.naver?where=nexearch&sm=top_hty&fbm=0"
    "&ie=utf8&query=%EB%B0%98%EB%8F%84%EC%B2%B4&ackey=lbkjibi4"
)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}


def fetch(url):
    """URL을 요청해 BeautifulSoup 객체를 반환한다."""
    res = requests.get(url, headers=HEADERS, timeout=10)
    res.raise_for_status()
    res.encoding = "utf-8"
    return BeautifulSoup(res.text, "html.parser")


def clean(text):
    """스크린리더용 '새 창 열림' 문구와 여분의 공백을 제거한다."""
    return re.sub(r"\s+", " ", text.replace("새 창 열림", "")).strip()


def gdid(anchor):
    """같은 기사에 속한 링크들이 공유하는 ID(data-nlog-params의 gdid)."""
    try:
        return json.loads(anchor.get("data-nlog-params", "{}")).get("gdid")
    except json.JSONDecodeError:
        return None


def by_gdid(soup, target):
    """data-heatmap-target 값이 target인 <a>를 gdid를 키로 하는 dict로 만든다."""
    return {gdid(a): a for a in soup.select(f"a[data-heatmap-target='{target}']") if gdid(a)}


def enclosing_profile(tag):
    """태그를 감싸는 가장 가까운 Profile 컴포넌트(언론사명 + 게시 시각)를 찾는다."""
    for parent in tag.parents:
        if parent.get("data-sds-comp") == "Profile":
            return parent
    return None


def parse_search(soup):
    """검색 결과 페이지에서 기사 목록(제목, 언론사, 시각, 요약, 링크, 썸네일)을 추출한다."""
    summaries = by_gdid(soup, ".body")
    images = by_gdid(soup, ".img")
    navs = by_gdid(soup, ".nav")  # 네이버뉴스 링크

    items = []
    for key, title_a in by_gdid(soup, ".tit").items():
        press, published = "", ""
        nav = navs.get(key)
        profile = enclosing_profile(nav) if nav else None
        if profile:
            name = profile.select_one(".sds-comps-profile-info-title-text")
            subtexts = [clean(t.get_text()) for t in profile.select(".sds-comps-profile-info-subtext")]
            press = clean(name.get_text()) if name else ""
            published = subtexts[0] if subtexts else ""

        body = summaries.get(key)
        img = images.get(key)
        img_tag = img.find("img") if img else None
        items.append({
            "title": clean(title_a.get_text()),
            "press": press,
            "published": published,
            "summary": clean(body.get_text()) if body else "",
            "press_url": title_a["href"],
            "naver_url": nav["href"] if nav else None,
            "thumbnail": img_tag.get("src", "") if img_tag else "",
        })
    return items


def parse_article(soup):
    """네이버 뉴스 기사 페이지(n.news.naver.com)에서 제목, 작성일, 본문을 추출한다."""
    title = soup.select_one("#title_area")
    body = soup.select_one("#dic_area")
    date = soup.select_one("span._ARTICLE_DATE_TIME[data-date-time]")

    if body:
        # 본문 안의 사진 설명, 스크립트 등은 제거
        for tag in body.select("script, style, .img_desc, em.img_desc, .end_photo_org"):
            tag.decompose()
    return {
        "title": clean(title.get_text()) if title else "",
        "date": date["data-date-time"] if date else "",
        "content": clean(body.get_text(" ")) if body else "",
    }


def build_result(item, article):
    """검색 결과 항목(item)과 기사 페이지 정보(article)를 하나의 결과 dict로 합친다."""
    return {
        "title": article.get("title") or item["title"],
        "press": item["press"],
        "published": item["published"],
        "date": article.get("date", ""),
        "url": item["naver_url"] or item["press_url"],
        "press_url": item["press_url"],
        "thumbnail": item["thumbnail"],
        "summary": item["summary"],
        "content": article.get("content", ""),
    }


def crawl(url, limit, delay):
    print(f"검색 결과 요청: {url}")
    items = parse_search(fetch(url))[:limit]
    print(f"기사 {len(items)}건 발견\n")

    results = []
    for i, item in enumerate(items, 1):
        print(f"[{i}/{len(items)}] {item['press']} | {item['title']} ({item['published']})")
        article = {}
        if item["naver_url"]:
            time.sleep(delay)  # 서버에 부담을 주지 않도록 요청 사이에 간격을 둔다
            try:
                article = parse_article(fetch(item["naver_url"]))
            except requests.RequestException as e:
                print(f"    본문 수집 실패: {e}")
        else:
            print("    네이버뉴스 링크가 없어 검색 결과의 요약만 저장합니다.")

        results.append(build_result(item, article))
    return results


EXCEL_COLUMNS = [
    # (머리글, 결과 dict의 키, 열 너비)
    ("언론사", "press", 14),
    ("제목", "title", 50),
    ("게시 시각", "published", 12),
    ("작성일", "date", 20),
    ("네이버 링크", "url", 40),
    ("원문 링크", "press_url", 40),
    ("요약", "summary", 50),
    ("본문", "content", 100),
]
_ILLEGAL_XML_CHARS = re.compile("[" + "".join(map(chr, [*range(0, 9), 11, 12, *range(14, 32)])) + "]")
EXCEL_CELL_LIMIT = 32767  # 엑셀 셀 하나에 넣을 수 있는 최대 글자 수


def save_excel(results, path):
    """수집 결과를 엑셀(.xlsx) 파일로 저장한다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"

    ws.append([name for name, _, _ in EXCEL_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in results:
        # 엑셀이 허용하지 않는 제어 문자를 제거하고, 셀 최대 길이를 넘지 않게 자른다
        ws.append([_ILLEGAL_XML_CHARS.sub("", r.get(key, ""))[:EXCEL_CELL_LIMIT] for _, key, _ in EXCEL_COLUMNS])

    link_cols = [i for i, (_, key, _) in enumerate(EXCEL_COLUMNS, 1) if key in ("url", "press_url")]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column in link_cols and cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    for i, (_, _, width) in enumerate(EXCEL_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"  # 머리글 고정
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def save_results(results, path):
    """확장자가 .xlsx이면 엑셀로, 그 외에는 JSON으로 저장한다."""
    if path.lower().endswith(".xlsx"):
        save_excel(results, path)
    else:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False, indent=2)


def main():
    parser = argparse.ArgumentParser(description="네이버 뉴스 검색 결과 크롤러")
    parser.add_argument("--url", default=DEFAULT_URL, help="네이버 검색 결과 URL (기본: '반도체' 검색)")
    parser.add_argument("--limit", type=int, default=10, help="수집할 최대 기사 수 (기본 10)")
    parser.add_argument("--delay", type=float, default=1.0, help="기사 요청 사이 대기 시간(초)")
    parser.add_argument("--out", default="news.json", help="저장할 파일 경로 (.json 또는 .xlsx)")
    args = parser.parse_args()

    try:
        results = crawl(args.url, args.limit, args.delay)
    except requests.RequestException as e:
        sys.exit(f"요청 실패: {e}")

    try:
        save_results(results, args.out)
    except OSError as e:
        sys.exit(f"저장 실패: {e}")

    print(f"\n{len(results)}건을 {args.out}에 저장했습니다.")
    for r in results[:3]:
        print(f"\n■ {r['title']} ({r['press']} {r['date'] or r['published']})")
        print(f"  {(r['content'] or r['summary'])[:150]}...")


if __name__ == "__main__":
    main()
